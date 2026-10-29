import os
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pymongo
from bson import ObjectId
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parents[1]
load_dotenv(BASE_DIR / ".env")

MONGO_URI = os.getenv("MONGO_URI")
DATABASE_NAME = os.getenv("MONGO_DATABASE", "heavycult")
OUTPUT_FILE = BASE_DIR / f"Catalogo_Datos_VentasDB_{datetime.now():%Y%m%d}.xlsx"
SAMPLE_SIZE = 1000

ROLE_BY_COLLECTION = {
    "customers": "OLTP",
    "products": "OLTP",
    "orders": "OLTP",
    "whatsapp_logs": "OLTP",
    "ventas_analiticas": "OLAP",
    "dim_cliente": "Data Warehouse - Dimension",
    "dim_producto": "Data Warehouse - Dimension",
    "dim_fecha": "Data Warehouse - Dimension",
    "fact_ventas": "Data Warehouse - Fact",
}

BUSINESS_DESCRIPTIONS = {
    "_id": "Identificador interno generado por MongoDB.",
    "phone": "Telefono unico del cliente.",
    "name": "Nombre de la persona o producto.",
    "city": "Ciudad asociada al cliente o envio.",
    "department": "Departamento geografico.",
    "created_at": "Fecha y hora de creacion del registro.",
    "order_number": "Numero unico de la orden.",
    "customer_id": "Referencia al cliente que realizo la orden.",
    "product_id": "Referencia al producto vendido.",
    "total_amount": "Valor total de la orden.",
    "total_venta": "Valor de venta usado en analitica.",
    "order_profit": "Ganancia neta estimada de la orden.",
    "ganancia_neta": "Ganancia neta usada en analitica.",
    "status": "Estado operativo de la orden.",
    "payment_method": "Metodo de pago utilizado.",
    "cliente_key": "Clave de la dimension de clientes.",
    "producto_key": "Clave de la dimension de productos.",
    "fecha_key": "Clave de la dimension calendario con formato YYYYMMDD.",
}


def clasificar_coleccion(nombre):
    return ROLE_BY_COLLECTION.get(nombre, "No clasificada")


def tipo_estandar(valor):
    if valor is None:
        return "Null"
    if isinstance(valor, bool):
        return "Boolean"
    if isinstance(valor, ObjectId):
        return "ObjectId"
    if isinstance(valor, datetime):
        return "Timestamp"
    if isinstance(valor, date):
        return "Date"
    if isinstance(valor, str):
        return "String"
    if isinstance(valor, int) and not isinstance(valor, bool):
        return "Integer"
    if isinstance(valor, (float, Decimal)):
        return "Decimal/Double"
    if isinstance(valor, list):
        return "Array"
    if isinstance(valor, dict):
        return "Object"
    return type(valor).__name__


def aplanar_documento(valor, ruta=""):
    campos = []
    if isinstance(valor, dict):
        for clave, elemento in valor.items():
            nueva_ruta = f"{ruta}.{clave}" if ruta else clave
            campos.extend(aplanar_documento(elemento, nueva_ruta))
    elif isinstance(valor, list):
        if not valor:
            campos.append((ruta, "Array"))
        elif all(isinstance(elemento, dict) for elemento in valor):
            campos.append((ruta, "Array"))
            for elemento in valor:
                for campo, tipo in aplanar_documento(elemento, f"{ruta}[]"):
                    campos.append((campo, tipo))
        else:
            campos.append((ruta, "Array"))
            for elemento in valor:
                campos.append((ruta, tipo_estandar(elemento)))
    else:
        campos.append((ruta, tipo_estandar(valor)))
    return campos


def es_fk(campo):
    ultimo = campo.split(".")[-1].replace("[]", "")
    return ultimo.endswith("_id") or ultimo in {"cliente_key", "producto_key", "fecha_key"}


def indice_por_campo(indices):
    resultado = defaultdict(list)
    for indice in indices:
        nombre = indice["name"]
        for campo, direccion in indice["key"].items():
            resultado[campo].append(f"{nombre} ({direccion})")
    return resultado


def construir_diccionario(db, colecciones):
    filas = []
    for nombre in colecciones:
        muestras = list(db[nombre].find({}).limit(SAMPLE_SIZE))
        tipos = defaultdict(set)
        for documento in muestras:
            for campo, tipo in aplanar_documento(documento):
                tipos[campo].add(tipo)

        indices = indice_por_campo(list(db[nombre].list_indexes()))
        for campo in sorted(tipos):
            ultimo = campo.split(".")[-1].replace("[]", "")
            pk = "SI" if campo == "_id" or ultimo.endswith("_key") else "NO"
            fk = "SI" if es_fk(campo) and pk == "NO" else "NO"
            filas.append({
                "Coleccion": nombre,
                "Rol": clasificar_coleccion(nombre),
                "Campo": campo,
                "Tipo de dato": " | ".join(sorted(tipos[campo])),
                "PK": pk,
                "FK": fk,
                "Indexado": "SI" if campo in indices else "NO",
                "Indices": "; ".join(indices.get(campo, [])),
                "Descripcion de negocio": BUSINESS_DESCRIPTIONS.get(ultimo, f"Campo {ultimo} de {nombre}."),
            })
    return filas


def construir_resumen(db, colecciones):
    filas = []
    for nombre in colecciones:
        estadisticas = db.command("collStats", nombre)
        filas.append({
            "Coleccion": nombre,
            "Rol arquitectonico": clasificar_coleccion(nombre),
            "Documentos": db[nombre].count_documents({}),
            "Indices": len(list(db[nombre].list_indexes())),
            "Tamano almacenamiento (bytes)": estadisticas.get("storageSize", 0),
            "Tamano datos (bytes)": estadisticas.get("size", 0),
        })
    return filas


def escribir_hoja(ws, titulo, filas):
    ws.title = titulo
    if not filas:
        ws.append(["Sin datos"])
        return
    encabezados = list(filas[0])
    ws.append(encabezados)
    for fila in filas:
        ws.append([fila[encabezado] for encabezado in encabezados])
    encabezado_fill = PatternFill("solid", fgColor="1F4E78")
    for celda in ws[1]:
        celda.fill = encabezado_fill
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for columna in ws.columns:
        maximo = max(len(str(celda.value or "")) for celda in columna)
        ws.column_dimensions[get_column_letter(columna[0].column)].width = min(max(maximo + 2, 12), 55)


def generar_catalogo():
    if not MONGO_URI:
        raise RuntimeError("MONGO_URI no esta configurado en .env.")

    client = pymongo.MongoClient(MONGO_URI, serverSelectionTimeoutMS=10000)
    try:
        client.admin.command("ping")
        db = client[DATABASE_NAME]
        colecciones = sorted(db.list_collection_names())
        resumen = construir_resumen(db, colecciones)
        diccionario = construir_diccionario(db, colecciones)

        libro = Workbook()
        escribir_hoja(libro.active, "Resumen Arquitectonico", resumen)
        escribir_hoja(libro.create_sheet(), "Diccionario de Datos", diccionario)
        libro.properties.title = "Catalogo de Datos - VentasDB"
        libro.properties.creator = "HeavyCult Data Platform"
        libro.save(OUTPUT_FILE)
        print(f"[CATALOGO] Colecciones inspeccionadas: {len(colecciones)}")
        print(f"[CATALOGO] Campos documentados: {len(diccionario)}")
        print(f"[CATALOGO] Archivo generado: {OUTPUT_FILE}")
    finally:
        client.close()


if __name__ == "__main__":
    generar_catalogo()